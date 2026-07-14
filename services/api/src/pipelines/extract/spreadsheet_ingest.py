from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Tuple

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None


IDENTITY_PATTERNS = ("producer", "domain", "ys", "estate", "brand", "name", "product", "wine", "label", "sku", "code")
PROFILE_PATTERNS = ("region", "country", "village", "vineyard", "style", "classification", "grape", "variet", "vintage", "appellation")
QUOTE_PATTERNS = ("price", "stock", "inventory", "qty", "quantity", "cost", "margin", "offer", "quote", "available")
SOURCE_PATTERNS = ("source", "critic", "score", "rating", "note", "notes", "url", "website", "reference")


def iter_spreadsheet_chunks(path: str | Path) -> Iterator[Dict[str, object]]:
    source_path = Path(path)
    extension = source_path.suffix.lower()
    if extension == ".csv":
        records = _iter_csv_rows(source_path)
    elif extension == ".xlsx":
        records = _iter_xlsx_rows(source_path)
    elif extension == ".xls":
        records = _iter_xls_rows(source_path)
    else:
        raise ValueError(f"Unsupported spreadsheet extension: {extension}")

    for sheet_name, row_number, row in records:
        normalized_row = {
            _normalize_header(key): _normalize_cell(value)
            for key, value in (row or {}).items()
            if _normalize_header(key) and _normalize_cell(value)
        }
        if not normalized_row:
            continue
        record_key = f"{source_path.name}:{sheet_name}:{row_number}"
        row_context = _extract_row_context(normalized_row)
        for group_name, group_fields in _group_row_fields(normalized_row):
            text = _render_group_text(source_path.name, sheet_name, row_number, record_key, group_name, group_fields)
            if not text:
                continue
            meta = {
                "sheet_name": sheet_name,
                "row_id": record_key,
                "row_number": row_number,
                "record_key": record_key,
                "source_family": "spreadsheet",
                "structured_group": group_name,
                "doc_type": _group_doc_type(group_name),
            }
            meta.update(row_context)
            yield {
                "text": text,
                "meta": meta,
            }


def _iter_csv_rows(path: Path) -> Iterator[Tuple[str, int, Dict[str, object]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp950"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                for index, row in enumerate(reader, start=2):
                    yield "Sheet1", index, row
            return
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    if last_error:
        raise last_error


def _iter_xlsx_rows(path: Path) -> Iterator[Tuple[str, int, Dict[str, object]]]:
    if load_workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl is required to ingest .xlsx files.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = _normalize_headers(next(rows, None))
        if not headers:
            continue
        for index, row in enumerate(rows, start=2):
            record = {headers[column]: row[column] for column in range(min(len(headers), len(row or []))) if headers[column]}
            yield sheet.title, index, record


def _iter_xls_rows(path: Path) -> Iterator[Tuple[str, int, Dict[str, object]]]:
    if xlrd is None:  # pragma: no cover
        raise RuntimeError("xlrd is required to ingest .xls files.")
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        if sheet.nrows <= 1:
            continue
        headers = _normalize_headers(sheet.row_values(0))
        if not headers:
            continue
        for row_index in range(1, sheet.nrows):
            values = sheet.row_values(row_index)
            record = {headers[column]: values[column] for column in range(min(len(headers), len(values))) if headers[column]}
            yield sheet.name, row_index + 1, record


def _normalize_headers(values: Iterable[object] | None) -> List[str]:
    headers: List[str] = []
    for value in values or []:
        headers.append(_normalize_header(value))
    return headers


def _normalize_header(value: object) -> str:
    header = " ".join(str(value or "").strip().split())
    if not header:
        return ""
    return header.lower()


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = " ".join(str(value).replace("\r", " ").replace("\n", " ").split())
    return text.strip()


def _group_row_fields(row: Dict[str, str]) -> Iterator[Tuple[str, Dict[str, str]]]:
    grouped = {
        "identity": {},
        "profile": {},
        "quote_inventory": {},
        "source_notes": {},
        "general": {},
    }
    for header, value in row.items():
        target = _classify_header(header)
        grouped[target][header] = value

    ordered_groups = ("identity", "profile", "quote_inventory", "source_notes", "general")
    for group_name in ordered_groups:
        fields = grouped[group_name]
        if not fields:
            continue
        yield group_name, fields


def _classify_header(header: str) -> str:
    lowered = header.lower()
    if _matches_any(lowered, IDENTITY_PATTERNS):
        return "identity"
    if _matches_any(lowered, PROFILE_PATTERNS):
        return "profile"
    if _matches_any(lowered, QUOTE_PATTERNS):
        return "quote_inventory"
    if _matches_any(lowered, SOURCE_PATTERNS):
        return "source_notes"
    return "general"


def _matches_any(header: str, patterns: Tuple[str, ...]) -> bool:
    return any(pattern in header for pattern in patterns)


def _render_group_text(
    filename: str,
    sheet_name: str,
    row_number: int,
    record_key: str,
    group_name: str,
    fields: Dict[str, str],
) -> str:
    if not fields:
        return ""
    ordered_items = list(fields.items())[:18]
    lines = [
        f"RAP spreadsheet source: {filename}",
        f"Sheet: {sheet_name}",
        f"Record: {record_key}",
        f"Row: {row_number}",
        f"Section: {group_name}",
    ]
    for header, value in ordered_items:
        cleaned_value = re.sub(r"\s+", " ", value).strip()
        if cleaned_value:
            lines.append(f"- {header}: {cleaned_value}")
    return "\n".join(lines).strip()


def _extract_row_context(row: Dict[str, str]) -> Dict[str, str]:
    producer = _pick_first_value(row, IDENTITY_PATTERNS)
    region = _pick_first_value(row, ("region", "country", "village", "appellation"))
    style = _pick_first_value(row, ("style", "classification", "grape", "variet"))
    wine_name = _pick_first_value(row, ("wine", "label", "product", "name"))
    vintage = _pick_first_value(row, ("vintage",))
    vineyard = _pick_first_value(row, ("vineyard",))
    page_url = _pick_first_value(row, ("url", "website", "reference", "source"))
    reviewer = _pick_first_value(row, ("reviewer", "critic", "author"))
    source_name = _pick_first_value(row, ("source name", "publication", "critic source", "website", "source"))
    score = _pick_first_value(row, ("score", "rating", "points"))
    issue_date = _pick_first_value(row, ("issue", "date", "published", "publication date"))
    page = _pick_first_value(row, ("page", "pdf page"))
    context: Dict[str, str] = {}
    if producer:
        context["producer"] = producer
    if region:
        context["region"] = region
    if style:
        context["style"] = style
    if wine_name:
        context["wine_name"] = wine_name
    if vintage:
        context["vintage"] = vintage
    if vineyard:
        context["vineyard"] = vineyard
    if page_url:
        context["page_url"] = page_url
    if reviewer:
        context["reviewer"] = reviewer
    if source_name:
        context["source_name"] = source_name
    if score:
        context["score"] = score
    if issue_date:
        context["issue_date"] = issue_date
    if page:
        context["page"] = page
    return context


def _pick_first_value(row: Dict[str, str], patterns: Tuple[str, ...]) -> str:
    for header, value in row.items():
        if value and _matches_any(header, patterns):
            return value
    return ""


def _group_doc_type(group_name: str) -> str:
    mapping = {
        "identity": "producer",
        "profile": "profile",
        "quote_inventory": "inventory",
        "source_notes": "source_note",
        "general": "reference",
    }
    return mapping.get(str(group_name or "").strip(), "reference")
