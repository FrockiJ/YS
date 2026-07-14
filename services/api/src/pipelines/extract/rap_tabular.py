from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None


HEADER_ALIASES = {
    "producer": "producer",
    "producer_name": "producer",
    "title_producer": "producer",
    "wine_name": "wine_name",
    "full_wine_name": "full_wine_name",
    "title_wine_full_name": "full_wine_name",
    "full_wine_name_": "full_wine_name",
    "vintage": "vintage",
    "region": "region",
    "full_region": "region",
    "wine_region": "region",
    "country": "country",
    "place_origin": "country",
    "appellation": "appellation",
    "classification": "classification",
    "vineyard_name": "vineyard_name",
    "vinyard_name_prev": "vineyard_name_previous",
    "wine_color": "wine_color",
    "wine_type": "wine_color",
    "color_type_sweetness": "wine_color",
    "grape_blend": "grape_blend",
    "grape_type": "grape_blend",
    "grape_variety": "grape_blend",
    "grape_details": "grape_blend",
    "point_score": "score_raw",
    "score": "score_raw",
    "rating": "score_raw",
    "star_point": "score_raw",
    "reviewer": "reviewer",
    "reviewed_by": "reviewer",
    "taster": "reviewer",
    "reviewer_title": "reviewer_title",
    "drinking_window": "drinking_window",
    "drink_window": "drinking_window",
    "drink_date": "drink_date",
    "full_tasting_note": "tasting_note",
    "tasting_note": "tasting_note",
    "producer_note": "producer_note",
    "published_date": "published_date",
    "review_date": "published_date",
    "tasting_date": "tasting_date",
    "release_price": "selling_price",
    "selling_price": "selling_price",
    "page_url": "page_url",
    "image_url": "image_url",
    "image_url1": "image_url",
    "title": "source_title",
    "cote": "subregion",
    "maturity": "maturity",
    "certified": "certified",
}

CANONICAL_HEADER_PREFERENCES = {
    "producer": ("producer", "producer_name", "title_producer"),
    "wine_name": ("wine_name",),
    "full_wine_name": ("full_wine_name", "title_wine_full_name", "full_wine_name_"),
    "vintage": ("vintage",),
    "region": ("region", "full_region", "wine_region"),
    "country": ("country", "place_origin"),
    "appellation": ("appellation",),
    "classification": ("classification",),
    "vineyard_name": ("vineyard_name", "vinyard_name_prev"),
    "wine_color": ("wine_color", "wine_type", "color_type_sweetness"),
    "grape_blend": ("grape_blend", "grape_type", "grape_variety", "grape_details"),
    "score_raw": ("point_score", "score", "rating", "star_point"),
    "reviewer": ("reviewer", "reviewed_by", "taster"),
    "reviewer_title": ("reviewer_title",),
    "drinking_window": ("drinking_window", "drink_window"),
    "drink_date": ("drink_date",),
    "tasting_note": ("full_tasting_note", "tasting_note"),
    "producer_note": ("producer_note",),
    "published_date": ("published_date", "review_date"),
    "tasting_date": ("tasting_date",),
    "selling_price": ("selling_price", "release_price"),
    "page_url": ("page_url",),
    "image_url": ("image_url", "image_url1"),
    "source_title": ("title",),
    "subregion": ("cote",),
    "maturity": ("maturity",),
    "certified": ("certified",),
}

REVIEW_REQUIRED_FIELDS = {"producer", "full_wine_name", "wine_name", "tasting_note", "score_raw", "page_url"}
CONTROL_WORKBOOK_NAMES = {"data scraping project.xlsx"}
CONTROL_TEXT_MARKERS = ("login", "site:", "password", "pw:", "scrape")


def iter_rap_tabular_payloads(path: str | Path) -> Iterator[Dict[str, object]]:
    source_path = Path(path)
    source_site = infer_source_site(source_path.name)
    for sheet_name, headers, rows in _iter_sheet_rows(source_path):
        sheet_kind = classify_sheet(source_path.name, sheet_name, headers, rows)
        if sheet_kind != "review_sheet":
            continue
        mapping = build_sheet_mapping(headers)
        for row_number, row in rows:
            normalized_row = {
                _normalize_header(key): _normalize_cell(value)
                for key, value in row.items()
                if _normalize_header(key) and _normalize_cell(value)
            }
            if not normalized_row:
                continue
            record = build_canonical_record(
                source_file=source_path.name,
                source_sheet=sheet_name,
                row_number=row_number,
                source_site=source_site,
                row=normalized_row,
                mapping=mapping,
            )
            if not _is_meaningful_record(record):
                continue
            yield {
                "record": record,
                "chunk_text": build_record_chunk_text(record),
                "chunk_meta": build_record_chunk_meta(record, sheet_name=sheet_name, row_number=row_number),
            }


def infer_source_site(filename: str) -> str:
    lowered = filename.lower()
    if "inside burgundy" in lowered:
        return "inside_burgundy"
    if "vinous" in lowered:
        return "vinous"
    if "wine advocate" in lowered or "robert parker" in lowered:
        return "wine_advocate"
    if "wine spectator" in lowered:
        return "wine_spectator"
    return "rap_tabular"


def classify_sheet(filename: str, sheet_name: str, headers: List[str], rows: List[Tuple[int, Dict[str, object]]]) -> str:
    lowered_file = filename.strip().lower()
    normalized_headers = [_normalize_header(header) for header in headers if _normalize_header(header)]
    preview_text = " ".join(
        _normalize_cell(value)
        for _, row in rows[:3]
        for value in row.values()
        if _normalize_cell(value)
    ).lower()

    if lowered_file in CONTROL_WORKBOOK_NAMES:
        return "instruction_sheet"
    if normalized_headers and all(header.startswith("http") for header in normalized_headers):
        return "url_seed_sheet"
    if len(normalized_headers) <= 1 and any(marker in preview_text for marker in CONTROL_TEXT_MARKERS):
        return "instruction_sheet"
    if len(normalized_headers) <= 1 and preview_text.startswith("http"):
        return "url_seed_sheet"
    mapped = {HEADER_ALIASES.get(header) for header in normalized_headers if HEADER_ALIASES.get(header)}
    if mapped & REVIEW_REQUIRED_FIELDS:
        return "review_sheet"
    return "instruction_sheet"


def build_sheet_mapping(headers: Iterable[str]) -> Dict[str, str]:
    grouped: Dict[str, List[str]] = {}
    for header in headers:
        normalized = _normalize_header(header)
        canonical = HEADER_ALIASES.get(normalized)
        if canonical:
            grouped.setdefault(canonical, []).append(normalized)
    mapping: Dict[str, str] = {}
    for canonical, candidates in grouped.items():
        mapping[canonical] = sorted(
            candidates,
            key=lambda item: _canonical_header_rank(canonical, item),
        )[0]
    return mapping


def build_canonical_record(
    *,
    source_file: str,
    source_sheet: str,
    row_number: int,
    source_site: str,
    row: Dict[str, str],
    mapping: Dict[str, str],
) -> Dict[str, object]:
    record_key = f"{source_file}:{source_sheet}:{row_number}"
    mapped_fields: Dict[str, object] = {}
    consumed_headers: set[str] = set()

    for canonical in mapping.keys():
        value, selected_header = _select_canonical_value(canonical, row=row, mapping=mapping)
        if not value or not selected_header:
            continue
        mapped_fields[canonical] = value
        consumed_headers.add(selected_header)

    image_urls = [
        value
        for header, value in row.items()
        if _normalize_header(header) in {"image_url", "image_url1"} and value
    ]
    if image_urls:
        mapped_fields["image_urls"] = image_urls
        consumed_headers.update({"image_url", "image_url1"})

    if "score_raw" in mapped_fields and "score_numeric" not in mapped_fields:
        mapped_fields["score_numeric"] = _extract_numeric_score(str(mapped_fields["score_raw"]))

    unmapped_fields = {
        header: value
        for header, value in row.items()
        if header not in consumed_headers and value
    }
    mapped_count = len(mapped_fields)
    total_count = max(1, len(row))
    mapping_confidence = round(mapped_count / total_count, 4)

    return {
        "record_key": record_key,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": row_number,
        "source_site": source_site,
        "producer": mapped_fields.get("producer"),
        "wine_name": mapped_fields.get("wine_name"),
        "full_wine_name": mapped_fields.get("full_wine_name"),
        "vintage": mapped_fields.get("vintage"),
        "region": mapped_fields.get("region"),
        "country": mapped_fields.get("country"),
        "appellation": mapped_fields.get("appellation"),
        "classification": mapped_fields.get("classification"),
        "vineyard_name": mapped_fields.get("vineyard_name"),
        "wine_color": mapped_fields.get("wine_color"),
        "grape_blend": mapped_fields.get("grape_blend"),
        "score_raw": mapped_fields.get("score_raw"),
        "score_numeric": mapped_fields.get("score_numeric"),
        "reviewer": mapped_fields.get("reviewer"),
        "reviewer_title": mapped_fields.get("reviewer_title"),
        "drinking_window": mapped_fields.get("drinking_window"),
        "drink_date": mapped_fields.get("drink_date"),
        "tasting_note": mapped_fields.get("tasting_note"),
        "producer_note": mapped_fields.get("producer_note"),
        "published_date": mapped_fields.get("published_date"),
        "tasting_date": mapped_fields.get("tasting_date"),
        "selling_price": mapped_fields.get("selling_price"),
        "page_url": mapped_fields.get("page_url"),
        "image_urls": mapped_fields.get("image_urls") or [],
        "subregion": mapped_fields.get("subregion"),
        "maturity": mapped_fields.get("maturity"),
        "certified": mapped_fields.get("certified"),
        "raw_row": row,
        "mapped_fields": mapped_fields,
        "unmapped_fields": unmapped_fields,
        "mapping_confidence": mapping_confidence,
        "schema_version": "rap_tabular_hybrid_v1",
    }


def _canonical_header_rank(canonical: str, header: str) -> tuple[int, str]:
    preferred = CANONICAL_HEADER_PREFERENCES.get(canonical, ())
    try:
        return (preferred.index(header), header)
    except ValueError:
        return (len(preferred) + 100, header)


def _select_canonical_value(
    canonical: str,
    *,
    row: Dict[str, str],
    mapping: Dict[str, str],
) -> Tuple[Optional[str], Optional[str]]:
    candidate_headers: List[str] = []
    primary_header = mapping.get(canonical)
    if primary_header:
        candidate_headers.append(primary_header)
    for header in row.keys():
        if HEADER_ALIASES.get(header) == canonical and header not in candidate_headers:
            candidate_headers.append(header)
    for header in sorted(candidate_headers, key=lambda item: _canonical_header_rank(canonical, item)):
        value = row.get(header)
        if value:
            return value, header
    return None, None


def build_record_chunk_text(record: Dict[str, object]) -> str:
    lines = [
        f"RAP tabular source: {record.get('source_site')}",
        f"Record: {record.get('record_key')}",
    ]
    for label, key in [
        ("Producer", "producer"),
        ("Wine", "full_wine_name"),
        ("Wine Name", "wine_name"),
        ("Vintage", "vintage"),
        ("Region", "region"),
        ("Country", "country"),
        ("Appellation", "appellation"),
        ("Classification", "classification"),
        ("Vineyard", "vineyard_name"),
        ("Color", "wine_color"),
        ("Grape", "grape_blend"),
        ("Score", "score_raw"),
        ("Reviewer", "reviewer"),
        ("Drinking Window", "drinking_window"),
        ("Published Date", "published_date"),
        ("Page URL", "page_url"),
    ]:
        value = record.get(key)
        if value:
            lines.append(f"{label}: {value}")
    if record.get("tasting_note"):
        lines.append(f"Tasting Note: {record['tasting_note']}")
    if record.get("producer_note"):
        lines.append(f"Producer Note: {record['producer_note']}")
    image_urls = record.get("image_urls") or []
    if image_urls:
        lines.append("Image URLs: " + ", ".join(str(url) for url in image_urls[:2]))
    return "\n".join(str(line).strip() for line in lines if str(line).strip())


def build_record_chunk_meta(
    record: Dict[str, object],
    *,
    sheet_name: Optional[str] = None,
    row_number: Optional[int] = None,
) -> Dict[str, object]:
    meta: Dict[str, object] = {
        "record_key": record.get("record_key"),
        "row_id": record.get("record_key"),
        "sheet_name": sheet_name or record.get("source_sheet"),
        "row_number": row_number or record.get("source_row"),
        "source_site": record.get("source_site"),
        "source_family": "licensed_review_dataset",
        "source_name": record.get("source_site") or record.get("source_name"),
        "reviewer": record.get("reviewer") or record.get("source_site"),
        "structured_group": "review_record",
        "doc_type": "source_note",
        "producer": record.get("producer"),
        "wine_name": record.get("wine_name"),
        "full_wine_name": record.get("full_wine_name"),
        "vintage": record.get("vintage"),
        "region": record.get("region"),
        "country": record.get("country"),
        "appellation": record.get("appellation"),
        "classification": record.get("classification"),
        "vineyard": record.get("vineyard_name"),
        "grape": record.get("grape_blend"),
        "color": record.get("wine_color"),
        "score_raw": record.get("score_raw"),
        "score": record.get("score_raw"),
        "issue_date": record.get("review_date") or record.get("issue_date") or record.get("date"),
        "page_url": record.get("page_url"),
        "mapping_confidence": record.get("mapping_confidence"),
    }
    source_trace = " | ".join(
        str(value)
        for value in (
            meta.get("source_name"),
            meta.get("producer"),
            meta.get("wine_name") or meta.get("full_wine_name"),
            meta.get("vintage"),
            f"row:{meta.get('row_id')}" if meta.get("row_id") else None,
            meta.get("page_url"),
        )
        if value
    )
    if source_trace:
        meta["source_trace"] = source_trace
    inferred_style = infer_record_style(record)
    if inferred_style:
        meta["style"] = inferred_style
    return {key: value for key, value in meta.items() if value not in (None, "", [], {})}


def infer_record_style(record: Dict[str, object]) -> Optional[str]:
    evidence = " ".join(
        str(record.get(key) or "")
        for key in (
            "region",
            "country",
            "appellation",
            "classification",
            "wine_name",
            "full_wine_name",
            "tasting_note",
            "producer_note",
            "page_url",
        )
    ).lower()
    sparkling_patterns = (
        "champagne",
        "sparkling",
        "mousseux",
        "pet nat",
        "pet-nat",
        "prosecco",
        "cava",
        "cremant",
        "crémant",
        "franciacorta",
    )
    still_patterns = (
        "still wine",
        "coteaux champenois",
        "tranquil",
    )
    if any(pattern in evidence for pattern in sparkling_patterns):
        return "sparkling"
    if any(pattern in evidence for pattern in still_patterns):
        return "still"
    return None


def _is_meaningful_record(record: Dict[str, object]) -> bool:
    return bool(
        record.get("producer")
        or record.get("full_wine_name")
        or record.get("wine_name")
        or record.get("tasting_note")
    )


def _extract_numeric_score(value: str) -> float | None:
    match = re.search(r"(\d{2,3}(?:\.\d+)?)", str(value or ""))
    if not match:
        return None
    try:
        return float(match.group(1))
    except ValueError:
        return None


def _iter_sheet_rows(path: Path) -> Iterator[Tuple[str, List[str], List[Tuple[int, Dict[str, object]]]]]:
    extension = path.suffix.lower()
    if extension == ".csv":
        yield from _iter_csv_sheet(path)
        return
    if extension == ".xlsx":
        yield from _iter_xlsx_sheets(path)
        return
    if extension == ".xls":
        yield from _iter_xls_sheets(path)
        return
    raise ValueError(f"Unsupported spreadsheet extension: {extension}")


def _iter_csv_sheet(path: Path) -> Iterator[Tuple[str, List[str], List[Tuple[int, Dict[str, object]]]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp950"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                headers = [str(name).strip() for name in (reader.fieldnames or []) if str(name or "").strip()]
                rows = [(index, row or {}) for index, row in enumerate(reader, start=2)]
                yield "Sheet1", headers, rows
            return
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    if last_error:
        raise last_error


def _iter_xlsx_sheets(path: Path) -> Iterator[Tuple[str, List[str], List[Tuple[int, Dict[str, object]]]]]:
    if load_workbook is None:  # pragma: no cover
        raise RuntimeError("openpyxl is required to ingest .xlsx files.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() for v in (next(rows, None) or []) if str(v or "").strip()]
            if not headers:
                continue
            materialized_rows: List[Tuple[int, Dict[str, object]]] = []
            for index, row in enumerate(rows, start=2):
                record = {
                    headers[column]: row[column]
                    for column in range(min(len(headers), len(row or [])))
                    if headers[column]
                }
                materialized_rows.append((index, record))
            yield sheet.title, headers, materialized_rows
    finally:
        workbook.close()


def _iter_xls_sheets(path: Path) -> Iterator[Tuple[str, List[str], List[Tuple[int, Dict[str, object]]]]]:
    if xlrd is None:  # pragma: no cover
        raise RuntimeError("xlrd is required to ingest .xls files.")
    workbook = xlrd.open_workbook(path)
    for sheet in workbook.sheets():
        if sheet.nrows <= 1:
            continue
        headers = [str(value).strip() for value in sheet.row_values(0) if str(value or "").strip()]
        if not headers:
            continue
        materialized_rows: List[Tuple[int, Dict[str, object]]] = []
        for row_index in range(1, sheet.nrows):
            values = sheet.row_values(row_index)
            record = {
                headers[column]: values[column]
                for column in range(min(len(headers), len(values)))
                if headers[column]
            }
            materialized_rows.append((row_index + 1, record))
        yield sheet.name, headers, materialized_rows


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return ""
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split()).strip()
