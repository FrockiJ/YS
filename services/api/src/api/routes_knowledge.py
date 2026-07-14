from __future__ import annotations

import csv
import hashlib
import html.parser
import io
import json
import re
import urllib.parse
import urllib.robotparser
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Optional

import httpx
import numpy as np
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pydantic import BaseModel, Field, HttpUrl

from ..core import config, db
from ..core.auth import require_permissions
from ..core.company_profile import COMPANY_PROFILE_FILENAME
from ..pipelines.extract.chunker import simple_chunks
from ..pipelines.nlp.embedder import embed_texts
from ..pipelines.rag.indexer import clear_document_payload, delete_documents_by_ids, insert_chunks, upsert_document

router = APIRouter(prefix="/knowledge", tags=["Knowledge"])

_SOURCE_RAG_TABLE = "rag_table"
_SOURCE_COMPANY_PROFILE = "company_profile"
_SOURCE_EXTERNAL_SITE = "external_site"
_COMPANY_PROFILE_FILENAME = COMPANY_PROFILE_FILENAME
_MAX_CHUNK_LEN = 1200
_USER_AGENT = "YS-KnowledgeIngest/1.0"


class CompanyProfilePayload(BaseModel):
    company_name: str = ""
    industry: str = ""
    product_categories: str = ""
    tone: str = ""
    service_scope: str = ""
    contact_info: str = ""
    faq: str = ""
    policies: str = ""
    language: str = "zh-TW"


class ExternalSitePayload(BaseModel):
    url: HttpUrl
    name: str = ""
    category: str = ""
    language: str = "zh-TW"
    trust_level: str = Field("approved", pattern="^(approved|reference|experimental)$")
    enabled: bool = True


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def _checksum_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _normalize_tags(value: str | None) -> List[str]:
    tags: List[str] = []
    for item in re.split(r"[,，\n]+", str(value or "")):
        tag = item.strip()
        if tag and tag not in tags:
            tags.append(tag)
    return tags


def _json_meta(value: Any) -> Dict[str, Any]:
    if isinstance(value, dict):
        return dict(value)
    if isinstance(value, str) and value.strip():
        try:
            parsed = json.loads(value)
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            return {}
    return {}


def _safe_filename(value: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value or "").strip()).strip("-")
    return normalized or "knowledge"


def _resize_embedding(values: Any) -> List[float]:
    try:
        vector = [float(item) for item in list(values)]
    except Exception:
        vector = []
    target_dim = int(getattr(config, "PGVECTOR_DIM", 384) or 384)
    if len(vector) == target_dim:
        return vector
    if len(vector) > target_dim:
        return vector[:target_dim]
    return vector + [0.0] * (target_dim - len(vector))


def _embed_or_zero(texts: List[str]) -> List[List[float]]:
    if not texts:
        return []
    try:
        vectors = embed_texts(texts)
        return [_resize_embedding(vector) for vector in vectors]
    except Exception:
        zero = _resize_embedding([])
        return [list(zero) for _ in texts]


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp950", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize_row(raw: Dict[str, Any]) -> Dict[str, str]:
    row: Dict[str, str] = {}
    for key, value in raw.items():
        clean_key = str(key or "").strip()
        if not clean_key:
            continue
        if value is None:
            clean_value = ""
        else:
            clean_value = str(value).strip()
        row[clean_key] = clean_value
    return row


def _parse_csv(data: bytes) -> List[Dict[str, str]]:
    text = _decode_csv(data)
    reader = csv.DictReader(io.StringIO(text))
    return [_normalize_row(row) for row in reader if any(str(v or "").strip() for v in row.values())]


def _parse_xlsx(data: bytes) -> List[Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise HTTPException(status_code=500, detail="openpyxl is required to import xlsx files") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    parsed: List[Dict[str, str]] = []
    for values in rows[1:]:
        raw = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
        row = _normalize_row(raw)
        if any(row.values()):
            parsed.append(row)
    return parsed


def _parse_table_file(filename: str, data: bytes) -> List[Dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _parse_csv(data)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(data)
    raise HTTPException(status_code=400, detail="Only csv and xlsx files are supported")


def _row_title(row: Dict[str, str], row_number: int) -> str:
    for key in ("title", "name", "topic", "question", "product", "product_name", "brand", "category"):
        value = row.get(key)
        if value:
            return value
    return f"Row {row_number}"


def _row_to_text(row: Dict[str, str], *, table_name: str, row_number: int) -> str:
    title = _row_title(row, row_number)
    lines = [f"Table: {table_name}", f"Record: {title}"]
    for key, value in row.items():
        if value:
            lines.append(f"{key}: {value}")
    return "\n".join(lines)


def _company_profile_text(payload: CompanyProfilePayload) -> str:
    fields = [
        ("Company / Brand", payload.company_name),
        ("Industry", payload.industry),
        ("Product categories", payload.product_categories),
        ("Preferred tone", payload.tone),
        ("Service scope", payload.service_scope),
        ("Contact info", payload.contact_info),
        ("FAQ", payload.faq),
        ("Policies", payload.policies),
    ]
    return "\n".join(f"{label}: {value}" for label, value in fields if str(value or "").strip())


def _build_chunks(
    texts: Iterable[str],
    *,
    source_family: str,
    source_name: str,
    base_meta: Dict[str, Any],
    row_metas: Optional[List[Dict[str, Any]]] = None,
) -> List[Dict[str, Any]]:
    chunk_texts: List[str] = []
    chunk_metas: List[Dict[str, Any]] = []
    for index, text in enumerate(texts):
        meta = dict(base_meta)
        if row_metas and index < len(row_metas):
            meta.update(row_metas[index])
        for part in simple_chunks(text, max_len=_MAX_CHUNK_LEN):
            chunk_texts.append(part)
            chunk_metas.append(meta)
    embeddings = _embed_or_zero(chunk_texts)
    chunks: List[Dict[str, Any]] = []
    for index, text in enumerate(chunk_texts):
        meta = dict(chunk_metas[index])
        meta.setdefault("source_family", source_family)
        meta.setdefault("source_name", source_name)
        meta.setdefault("allowed_use", ["rag_table" if source_family == _SOURCE_RAG_TABLE else "rag_knowledge"])
        meta.setdefault("source_tier", "internal_approved")
        chunks.append(
            {
                "chunk_idx": index,
                "text": text,
                "summary": text[:500],
                "keywords": [source_family, source_name],
                "meta": meta,
                "kg": {"nodes": [], "edges": []},
                "embedding": embeddings[index],
            }
        )
    return chunks


async def _list_documents(source_family: str) -> List[Dict[str, Any]]:
    conn = await db.get_conn()
    try:
        rows = await conn.fetch(
            """
            SELECT d.id, d.filename, d.filetype, d.meta, d.created_at, COUNT(c.id)::int AS chunk_count
              FROM documents d
              LEFT JOIN chunks c ON c.document_id = d.id
             WHERE d.meta->>'source_family' = $1
             GROUP BY d.id
             ORDER BY d.id DESC
            """,
            source_family,
        )
        items: List[Dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            item["meta"] = _json_meta(item.get("meta"))
            items.append(item)
        return items
    finally:
        await conn.close()


@router.get("/rag-tables", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def list_rag_tables() -> Dict[str, Any]:
    rows = await _list_documents(_SOURCE_RAG_TABLE)
    return {"items": rows}


@router.post("/rag-tables/import", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def import_rag_table(
    file: UploadFile = File(...),
    table_name: str = Form(""),
    language: str = Form("zh-TW"),
    source_name: str = Form(""),
    tags: str = Form(""),
    replace: bool = Form(True),
) -> Dict[str, Any]:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    rows = _parse_table_file(file.filename or "", data)
    if not rows:
        raise HTTPException(status_code=400, detail="No data rows found")
    resolved_table = table_name.strip() or source_name.strip() or file.filename or "RAG table"
    filename = f"rag-table:{_safe_filename(resolved_table)}"
    row_metas = [
        {
            "row_number": index,
            "raw_row": row,
            "table_name": resolved_table,
        }
        for index, row in enumerate(rows, start=2)
    ]
    meta = {
        "filetype": "rag_table",
        "source_family": _SOURCE_RAG_TABLE,
        "source_name": source_name.strip() or resolved_table,
        "table_name": resolved_table,
        "language": language,
        "tags": _normalize_tags(tags),
        "row_count": len(rows),
        "imported_at": _utcnow(),
    }
    chunks = _build_chunks(
        [_row_to_text(row, table_name=resolved_table, row_number=index) for index, row in enumerate(rows, start=1)],
        source_family=_SOURCE_RAG_TABLE,
        source_name=meta["source_name"],
        base_meta=meta,
        row_metas=row_metas,
    )
    conn = await db.get_conn()
    try:
        doc_id, doc_status = await upsert_document(
            conn,
            filename,
            meta,
            checksum=_checksum_bytes(data),
            filesize=len(data),
        )
        if replace:
            await clear_document_payload(conn, [doc_id], filenames=[filename])
        inserted = await insert_chunks(conn, doc_id, chunks)
        return {
            "document_id": doc_id,
            "status": doc_status,
            "table_name": resolved_table,
            "rows": len(rows),
            "chunks": len(inserted),
        }
    finally:
        await conn.close()


@router.delete("/rag-tables/{document_id}", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def delete_rag_table(document_id: int) -> Dict[str, Any]:
    conn = await db.get_conn()
    try:
        row = await conn.fetchrow(
            "SELECT id, filename FROM documents WHERE id=$1 AND meta->>'source_family'=$2",
            document_id,
            _SOURCE_RAG_TABLE,
        )
        if not row:
            raise HTTPException(status_code=404, detail="RAG table not found")
        result = await delete_documents_by_ids(conn, [document_id], filenames=[row["filename"]])
        return result
    finally:
        await conn.close()


@router.get("/company-profile", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def get_company_profile() -> Dict[str, Any]:
    conn = await db.get_conn()
    try:
        row = await conn.fetchrow(
            "SELECT id, filename, meta, created_at FROM documents WHERE filename=$1 LIMIT 1",
            _COMPANY_PROFILE_FILENAME,
        )
        if not row:
            return {"profile": {}, "document_id": None, "chunk_count": 0}
        count_row = await conn.fetchrow("SELECT COUNT(*)::int AS count FROM chunks WHERE document_id=$1", row["id"])
        meta = _json_meta(row["meta"])
        return {
            "profile": meta.get("profile") or {},
            "document_id": row["id"],
            "chunk_count": int((count_row or {}).get("count") or 0),
            "updated_at": meta.get("updated_at") or meta.get("imported_at"),
        }
    finally:
        await conn.close()


@router.post("/company-profile", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def save_company_profile(payload: CompanyProfilePayload) -> Dict[str, Any]:
    text = _company_profile_text(payload)
    if not text.strip():
        raise HTTPException(status_code=400, detail="Company profile is empty")
    profile = payload.model_dump()
    meta = {
        "filetype": "company_profile",
        "source_family": _SOURCE_COMPANY_PROFILE,
        "source_name": payload.company_name or "Company profile",
        "language": payload.language,
        "profile": profile,
        "updated_at": _utcnow(),
        "allowed_use": ["rag_company_profile", "rag_knowledge"],
        "source_tier": "internal_approved",
    }
    chunks = _build_chunks(
        [text],
        source_family=_SOURCE_COMPANY_PROFILE,
        source_name=meta["source_name"],
        base_meta=meta,
    )
    conn = await db.get_conn()
    try:
        doc_id, status_text = await upsert_document(
            conn,
            _COMPANY_PROFILE_FILENAME,
            meta,
            checksum=hashlib.sha256(json.dumps(profile, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
        )
        await clear_document_payload(conn, [doc_id], filenames=[_COMPANY_PROFILE_FILENAME])
        inserted = await insert_chunks(conn, doc_id, chunks)
        return {"document_id": doc_id, "status": status_text, "chunks": len(inserted), "profile": profile}
    finally:
        await conn.close()


@router.get("/external-sites", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def list_external_sites() -> Dict[str, Any]:
    rows = await _list_documents(_SOURCE_EXTERNAL_SITE)
    return {"items": rows}


@router.post("/external-sites", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def save_external_site(payload: ExternalSitePayload) -> Dict[str, Any]:
    meta = {
        "filetype": "external_site",
        "source_family": _SOURCE_EXTERNAL_SITE,
        "source_name": payload.name or str(payload.url),
        "url": str(payload.url),
        "category": payload.category,
        "language": payload.language,
        "trust_level": payload.trust_level,
        "enabled": payload.enabled,
        "updated_at": _utcnow(),
        "allowed_use": ["rag_external_site", "rag_knowledge"],
        "source_tier": "approved_external" if payload.trust_level == "approved" else payload.trust_level,
    }
    conn = await db.get_conn()
    try:
        doc_id, status_text = await upsert_document(
            conn,
            f"external-site:{str(payload.url)}",
            meta,
            checksum=hashlib.sha256(str(payload.url).encode("utf-8")).hexdigest(),
        )
        return {"document_id": doc_id, "status": status_text, "site": meta}
    finally:
        await conn.close()


class _TextExtractor(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.title = ""
        self.meta_description = ""
        self._capture_title = False
        self._skip_depth = 0
        self.paragraphs: List[str] = []

    def handle_starttag(self, tag: str, attrs: List[tuple[str, Optional[str]]]) -> None:
        lower = tag.lower()
        if lower in {"script", "style", "nav", "footer", "header"}:
            self._skip_depth += 1
        if lower == "title":
            self._capture_title = True
        if lower == "meta":
            attrs_dict = {str(key or "").lower(): str(value or "") for key, value in attrs}
            if attrs_dict.get("name", "").lower() == "description":
                self.meta_description = attrs_dict.get("content", "").strip()

    def handle_endtag(self, tag: str) -> None:
        lower = tag.lower()
        if lower in {"script", "style", "nav", "footer", "header"} and self._skip_depth:
            self._skip_depth -= 1
        if lower == "title":
            self._capture_title = False

    def handle_data(self, data: str) -> None:
        text = " ".join(str(data or "").split()).strip()
        if not text:
            return
        if self._capture_title:
            self.title = (self.title + " " + text).strip()
            return
        if self._skip_depth:
            return
        if len(text) >= 40:
            self.paragraphs.append(text)


async def _robots_allowed(url: str) -> bool:
    parsed = urllib.parse.urlparse(url)
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    parser = urllib.robotparser.RobotFileParser()
    try:
        async with httpx.AsyncClient(timeout=8, headers={"User-Agent": _USER_AGENT}) as client:
            response = await client.get(robots_url)
        if response.status_code >= 400:
            return True
        parser.parse(response.text.splitlines())
        return parser.can_fetch(_USER_AGENT, url)
    except Exception:
        return True


async def _fetch_site_text(url: str) -> Dict[str, Any]:
    if not await _robots_allowed(url):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="robots.txt does not allow ingesting this URL")
    async with httpx.AsyncClient(timeout=15, headers={"User-Agent": _USER_AGENT}) as client:
        response = await client.get(url, follow_redirects=True)
    if response.status_code >= 400:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=f"Unable to fetch URL: HTTP {response.status_code}")
    extractor = _TextExtractor()
    extractor.feed(response.text)
    paragraphs = []
    seen = set()
    for text in [extractor.meta_description, *extractor.paragraphs]:
        cleaned = " ".join(str(text or "").split()).strip()
        if cleaned and cleaned not in seen:
            paragraphs.append(cleaned)
            seen.add(cleaned)
        if len(paragraphs) >= 24:
            break
    if not paragraphs:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="No readable page content was found")
    return {"title": extractor.title, "paragraphs": paragraphs, "final_url": str(response.url)}


@router.post("/external-sites/{document_id}/ingest", dependencies=[Depends(require_permissions("admin.settings.access"))])
async def ingest_external_site(document_id: int) -> Dict[str, Any]:
    conn = await db.get_conn()
    try:
        row = await conn.fetchrow(
            "SELECT id, filename, meta FROM documents WHERE id=$1 AND meta->>'source_family'=$2",
            document_id,
            _SOURCE_EXTERNAL_SITE,
        )
        if not row:
            raise HTTPException(status_code=404, detail="External site not found")
        meta = _json_meta(row["meta"])
        if not bool(meta.get("enabled", True)):
            raise HTTPException(status_code=409, detail="External site is disabled")
        url = str(meta.get("url") or "").strip()
        if not url:
            raise HTTPException(status_code=400, detail="External site URL is missing")
        fetched = await _fetch_site_text(url)
        source_name = meta.get("source_name") or fetched.get("title") or url
        ingest_meta = {
            **meta,
            "source_name": source_name,
            "title": fetched.get("title") or source_name,
            "final_url": fetched.get("final_url") or url,
            "ingested_at": _utcnow(),
        }
        chunks = _build_chunks(
            fetched["paragraphs"],
            source_family=_SOURCE_EXTERNAL_SITE,
            source_name=source_name,
            base_meta=ingest_meta,
        )
        await upsert_document(conn, row["filename"], ingest_meta, checksum=hashlib.sha256(url.encode("utf-8")).hexdigest())
        await clear_document_payload(conn, [document_id], filenames=[row["filename"]])
        inserted = await insert_chunks(conn, document_id, chunks)
        return {"document_id": document_id, "chunks": len(inserted), "title": fetched.get("title"), "url": url}
    finally:
        await conn.close()
