import re
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from ..utils.cerp_export_lookup import (
    find_export_row_by_code,
    search_export_rows,
)
from .cerp_service import CerpService
from .label_print_service import LabelPrintService


class LabelPrintImportService:
    TEMPLATE_HEADERS = [
        "產品編號",
        "品名",
        "品牌",
        "規格",
        "價格",
        "大尺寸",
        "中尺寸",
        "小尺寸",
    ]
    SIZE_COLUMNS = [
        ("大尺寸", "large", "大尺寸"),
        ("中尺寸", "medium", "中尺寸"),
        ("小尺寸", "small", "小尺寸"),
    ]
    CHECKED_VALUES = {"y", "1", "true", "是"}
    _PRICE_PATTERN = re.compile(r"^(?:\d+|\d{1,3}(?:,\d{3})+)$")
    _NORMALIZE_RE = re.compile(r"[^\w]+", re.UNICODE)

    def __init__(
        self,
        *,
        label_print_service: Optional[LabelPrintService] = None,
        cerp_service: Optional[CerpService] = None,
    ) -> None:
        self._label_print_service = label_print_service or LabelPrintService()
        self._cerp_service = cerp_service or CerpService()

    async def import_workbook(self, user_id: int, file_bytes: bytes, lang: str = "zh-TW") -> Dict[str, Any]:
        workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
        if not workbook.sheetnames:
            raise ValueError("Workbook does not contain any worksheets")

        sheet = workbook[workbook.sheetnames[0]]
        headers = self._read_headers(sheet)
        if headers != self.TEMPLATE_HEADERS:
            raise ValueError("Invalid import template headers")

        errors: List[Dict[str, Any]] = []
        imported_rows = 0
        created_items = 0
        merged_rows = 0
        added_quantity = 0
        total_rows = 0

        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(self.TEMPLATE_HEADERS), values_only=True),
            start=2,
        ):
            raw_row = self._build_raw_row(values)
            if self._is_blank_row(raw_row):
                continue

            total_rows += 1
            try:
                normalized = self._normalize_row(raw_row)
                product = self._resolve_product(normalized)
            except RowImportError as exc:
                errors.append(self._build_error_payload(row_number, exc, raw_row))
                continue

            row_has_success = False
            for size_entry in normalized["sizes"]:
                add_result = await self._label_print_service.add_item(
                    user_id,
                    product["code"],
                    size_entry["size"],
                    product_snapshot=product["product"],
                    quantity=1,
                    price_override=normalized["price_override"],
                )
                status = add_result.get("status")
                if status == "invalid":
                    errors.append(
                        self._build_error_payload(
                            row_number,
                            RowImportError("invalid_row", f"{size_entry['label']}資料無效"),
                            raw_row,
                        )
                    )
                    continue
                if status == "conflict":
                    errors.append(
                        self._build_error_payload(
                            row_number,
                            RowImportError(
                                "price_conflict",
                                f"{size_entry['label']}價格與既有列印清單項目衝突",
                            ),
                            raw_row,
                        )
                    )
                    continue

                row_has_success = True
                added_quantity += 1
                if status == "created":
                    created_items += 1
                else:
                    merged_rows += 1

            if row_has_success:
                imported_rows += 1

        total_quantity = await self._label_print_service.get_total_quantity(user_id)
        return {
            "lang": lang,
            "total_rows": total_rows,
            "imported_rows": imported_rows,
            "created_items": created_items,
            "merged_rows": merged_rows,
            "added_quantity": added_quantity,
            "error_rows": errors,
            "total_quantity": total_quantity,
        }

    @classmethod
    def _build_error_payload(
        cls,
        row_number: int,
        error: "RowImportError",
        raw_row: Dict[str, Any],
    ) -> Dict[str, Any]:
        return {
            "row_number": row_number,
            "reason_code": error.reason_code,
            "message": error.message,
            "raw_row": raw_row,
        }

    @classmethod
    def _read_headers(cls, sheet) -> List[str]:
        return [str(cell.value or "").strip() for cell in sheet[1][: len(cls.TEMPLATE_HEADERS)]]

    @classmethod
    def _build_raw_row(cls, values: Any) -> Dict[str, Any]:
        cells = list(values or [])
        padded = cells + [None] * max(0, len(cls.TEMPLATE_HEADERS) - len(cells))
        return {header: padded[index] for index, header in enumerate(cls.TEMPLATE_HEADERS)}

    @classmethod
    def _is_blank_row(cls, raw_row: Dict[str, Any]) -> bool:
        return all(cls._is_blank_value(value) for value in raw_row.values())

    @staticmethod
    def _is_blank_value(value: Any) -> bool:
        if value is None or value is False:
            return True
        if isinstance(value, str) and not value.strip():
            return True
        return False

    @classmethod
    def _normalize_row(cls, raw_row: Dict[str, Any]) -> Dict[str, Any]:
        code = cls._normalize_text(raw_row.get("產品編號")).upper()
        name = cls._normalize_text(raw_row.get("品名"))
        brand = cls._normalize_text(raw_row.get("品牌"))
        specification = cls._normalize_text(raw_row.get("規格"))
        if not any([code, name]):
            raise RowImportError("missing_identifier", "至少需要填寫 產品編號、品名 其中之一")

        price_override = cls._normalize_price(raw_row.get("價格"))
        if raw_row.get("價格") not in (None, "") and price_override is None:
            raise RowImportError("invalid_price", "價格只能是正整數，可使用千分位逗號")

        sizes = cls._normalize_sizes(raw_row)
        if not sizes:
            raise RowImportError("missing_size_selection", "至少需要勾選一個尺寸欄位")

        return {
            "code": code,
            "name": name,
            "brand": brand,
            "specification": specification,
            "price_override": price_override,
            "sizes": sizes,
        }

    @classmethod
    def _normalize_sizes(cls, raw_row: Dict[str, Any]) -> List[Dict[str, str]]:
        sizes: List[Dict[str, str]] = []
        for column_name, size_key, size_label in cls.SIZE_COLUMNS:
            if cls._is_checked_value(raw_row.get(column_name)):
                sizes.append({"column": column_name, "size": size_key, "label": size_label})
        return sizes

    @classmethod
    def _is_checked_value(cls, value: Any) -> bool:
        if value is True:
            return True
        if value in (None, "", False):
            return False
        if isinstance(value, (int, float)):
            return float(value) == 1.0
        return str(value).strip().lower() in cls.CHECKED_VALUES

    @classmethod
    def _normalize_price(cls, value: Any) -> Optional[int]:
        if value in (None, ""):
            return None
        if isinstance(value, float):
            if not value.is_integer():
                return None
            return int(value) if int(value) > 0 else None
        if isinstance(value, int):
            return value if value > 0 else None
        text = str(value).strip()
        if not text or not cls._PRICE_PATTERN.fullmatch(text):
            return None
        number = int(text.replace(",", ""))
        return number if number > 0 else None

    @staticmethod
    def _normalize_text(value: Any) -> str:
        return str(value or "").strip()

    @classmethod
    def _normalize_label(cls, value: Any) -> str:
        normalized = cls._NORMALIZE_RE.sub(" ", str(value or "").lower())
        return re.sub(r"\s+", " ", normalized).strip()

    def _resolve_product(self, normalized_row: Dict[str, Any]) -> Dict[str, Any]:
        code = normalized_row["code"]
        if code:
            row = find_export_row_by_code(code)
            if not row:
                raise RowImportError("code_not_found", f"找不到產品編號 {code}")
            return self._build_product_payload(row)
        return self._resolve_by_name(normalized_row)

    def _resolve_by_name(self, normalized_row: Dict[str, Any]) -> Dict[str, Any]:
        name = normalized_row["name"]
        candidates = search_export_rows(name, limit=50)
        if not candidates:
            raise RowImportError("not_found", f"找不到品名 {name}")

        exact_name_matches = self._filter_exact_name_matches(candidates, name)
        if exact_name_matches:
            candidates = exact_name_matches
        if normalized_row["brand"]:
            candidates = self._filter_by_brand(candidates, normalized_row["brand"])
        if normalized_row["specification"]:
            candidates = self._filter_by_specification(candidates, normalized_row["specification"])

        unique_candidates = self._dedupe_candidates(candidates)
        if not unique_candidates:
            raise RowImportError("not_found", f"找不到符合條件的產品：{name}")
        if len(unique_candidates) > 1:
            raise RowImportError(
                "ambiguous_match",
                f"品名 {name} 對應到多筆產品，請補產品編號、品牌或版本",
            )
        return self._build_product_payload(unique_candidates[0])

    def _build_product_payload(self, row: Dict[str, Any]) -> Dict[str, Any]:
        code = str(row.get("invn002") or "").strip().upper()
        product = self._cerp_service.map_row_to_product(row, row.get("wd4inv1as") or [])
        return {"code": code, "product": product}

    @classmethod
    def _dedupe_candidates(cls, candidates: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        by_code: Dict[str, Dict[str, Any]] = {}
        for row in candidates:
            code = str(row.get("invn002") or "").strip().upper()
            if code:
                by_code[code] = row
        return list(by_code.values())

    @classmethod
    def _filter_exact_name_matches(cls, candidates: List[Dict[str, Any]], name: str) -> List[Dict[str, Any]]:
        normalized_name = cls._normalize_label(name)
        exact = []
        for row in candidates:
            name_candidates = [row.get("invn005"), row.get("invn077")]
            if any(cls._normalize_label(value) == normalized_name for value in name_candidates):
                exact.append(row)
        return cls._dedupe_candidates(exact)

    @classmethod
    def _filter_by_brand(cls, candidates: List[Dict[str, Any]], brand: str) -> List[Dict[str, Any]]:
        normalized_brand = cls._normalize_label(brand)
        matched = [
            row for row in candidates if cls._normalize_label(row.get("invn006")) == normalized_brand
        ]
        return cls._dedupe_candidates(matched)

    @staticmethod
    def _filter_by_specification(candidates: List[Dict[str, Any]], specification: str) -> List[Dict[str, Any]]:
        target = str(specification or "").strip()
        matched = []
        for row in candidates:
            candidate_values = [
                str(row.get("invn051") or "").strip(),
                str(row.get("invn807") or "").strip(),
            ]
            if target in candidate_values:
                matched.append(row)
        return matched


class RowImportError(Exception):
    def __init__(self, reason_code: str, message: str) -> None:
        super().__init__(message)
        self.reason_code = reason_code
        self.message = message
