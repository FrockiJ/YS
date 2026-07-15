from __future__ import annotations

import json
import zlib
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..core.config import YS_FAKE_CERP_DEFAULT_XLSX
from ..core.database import SessionLocal
from ..models.fake_cerp_product import FakeCerpProduct
from .fake_cerp_pricing import resolve_fake_list_price


class FakeCerpImportService:
    REQUIRED_HEADERS = {"code", "name"}
    OPTIONAL_HEADERS = {"spec1", "amount", "stock"}

    @staticmethod
    def deterministic_stock(code: str) -> int:
        normalized = str(code or "").strip().upper()
        if not normalized:
            return 0
        return zlib.crc32(normalized.encode("utf-8")) % 51

    @staticmethod
    def resolve_default_workbook_path() -> Path:
        raw = Path(str(YS_FAKE_CERP_DEFAULT_XLSX or "")).expanduser()
        if raw.is_absolute():
            return raw
        for parent in Path(__file__).resolve().parents:
            candidate = parent / raw
            if candidate.exists():
                return candidate
        return Path.cwd() / raw

    @staticmethod
    def resolve_camping_fixture_path() -> Path:
        for parent in Path(__file__).resolve().parents:
            candidate = parent / "cerp" / "fake_camping_products.json"
            if candidate.exists():
                return candidate
        return Path.cwd() / "cerp" / "fake_camping_products.json"

    @classmethod
    def _normalize_header(cls, value: Any) -> str:
        return str(value or "").strip().lower()

    @classmethod
    def _normalize_text(cls, value: Any) -> str:
        return str(value or "").strip()

    @classmethod
    def _normalize_decimal(cls, value: Any) -> Decimal:
        if value in (None, ""):
            return Decimal("0")
        try:
            return Decimal(str(value).replace(",", "").strip() or "0")
        except (InvalidOperation, ValueError):
            return Decimal("0")

    @classmethod
    def _normalize_stock(cls, value: Any, code: str) -> int:
        if value not in (None, ""):
            try:
                return max(int(Decimal(str(value).replace(",", "").strip())), 0)
            except (InvalidOperation, ValueError):
                pass
        return cls.deterministic_stock(code)

    @classmethod
    def _resolve_list_price(cls, value: Any, *, code: str, name: str, spec1: str = "") -> Decimal:
        return resolve_fake_list_price(value, code=code, name=name, spec1=spec1)

    @classmethod
    def _json_safe_value(cls, value: Any) -> Any:
        if isinstance(value, (str, int, float, bool)) or value is None:
            return value
        if isinstance(value, Decimal):
            return str(value)
        if isinstance(value, datetime):
            return value.isoformat()
        return str(value)

    @classmethod
    def _iter_rows_from_workbook(cls, file_bytes: bytes) -> Iterable[tuple[int, Dict[str, Any]]]:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
        if not workbook.sheetnames:
            raise ValueError("Workbook does not contain any worksheets")
        sheet = workbook[workbook.sheetnames[0]]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [cls._normalize_header(cell) for cell in (header_row or [])]
        if not cls.REQUIRED_HEADERS.issubset(set(headers)):
            raise ValueError("Fake CERP import requires headers: code, name")
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            raw = {
                headers[index]: cls._json_safe_value(values[index] if index < len(values) else None)
                for index in range(len(headers))
                if headers[index]
            }
            yield row_number, raw

    @classmethod
    def _normalize_import_rows(cls, raw_rows: Iterable[tuple[int, Dict[str, Any]]]) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        rows: List[Dict[str, Any]] = []
        errors: List[Dict[str, Any]] = []
        for row_number, raw in raw_rows:
            if not any(str(value or "").strip() for value in raw.values()):
                continue
            code = cls._normalize_text(raw.get("code")).upper()
            name = cls._normalize_text(raw.get("name"))
            if not code or not name:
                errors.append(
                    {
                        "row_number": row_number,
                        "reason_code": "missing_required_field",
                        "message": "code and name are required",
                        "raw_row": raw,
                    }
                )
                continue
            rows.append(
                {
                    "row_number": row_number,
                    "code": code,
                    "name": name,
                    "spec1": cls._normalize_text(raw.get("spec1")),
                    "amount": cls._resolve_list_price(
                        raw.get("amount"),
                        code=code,
                        name=name,
                        spec1=cls._normalize_text(raw.get("spec1")),
                    ),
                    "stock": cls._normalize_stock(raw.get("stock"), code),
                    "raw_row": raw,
                }
            )
        return rows, errors

    @classmethod
    def parse_workbook(cls, file_bytes: bytes) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        return cls._normalize_import_rows(cls._iter_rows_from_workbook(file_bytes))

    @classmethod
    def _iter_rows_from_json(cls, payload: Any) -> Iterable[tuple[int, Dict[str, Any]]]:
        data = json.loads(payload.decode("utf-8")) if isinstance(payload, (bytes, bytearray)) else payload
        if isinstance(data, dict):
            rows = data.get("products") or data.get("items") or data.get("rows")
        else:
            rows = data
        if not isinstance(rows, list):
            raise ValueError("Fake CERP JSON import requires a list or an object with products/items/rows")
        for row_number, raw in enumerate(rows, start=1):
            if not isinstance(raw, dict):
                yield row_number, {"code": "", "name": "", "_invalid_row": cls._json_safe_value(raw)}
                continue
            yield row_number, {str(key).strip(): cls._json_safe_value(value) for key, value in raw.items() if str(key).strip()}

    @classmethod
    def parse_json_rows(cls, payload: Any) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        return cls._normalize_import_rows(cls._iter_rows_from_json(payload))

    async def import_workbook(
        self,
        file_bytes: bytes,
        *,
        session: Optional[AsyncSession] = None,
    ) -> Dict[str, Any]:
        rows, errors = self.parse_workbook(file_bytes)
        return await self.import_rows(rows, errors=errors, session=session)

    async def import_json_rows(
        self,
        payload: Any,
        *,
        session: Optional[AsyncSession] = None,
    ) -> Dict[str, Any]:
        rows, errors = self.parse_json_rows(payload)
        return await self.import_rows(rows, errors=errors, session=session)

    async def import_rows(
        self,
        rows: List[Dict[str, Any]],
        *,
        errors: Optional[List[Dict[str, Any]]] = None,
        session: Optional[AsyncSession] = None,
    ) -> Dict[str, Any]:
        errors = errors or []
        owns_session = session is None
        db = session or SessionLocal()
        try:
            created = 0
            updated = 0
            now = datetime.now(timezone.utc)
            for row in rows:
                result = await db.execute(select(FakeCerpProduct).where(FakeCerpProduct.code == row["code"]))
                product = result.scalar_one_or_none()
                if product is None:
                    product = FakeCerpProduct(code=row["code"])
                    db.add(product)
                    created += 1
                else:
                    updated += 1
                product.name = row["name"]
                product.spec1 = row["spec1"] or None
                product.amount = row["amount"]
                product.stock = row["stock"]
                product.raw_row = row["raw_row"]
                product.imported_at = now
                product.updated_at = now
            if owns_session:
                await db.commit()
            return {
                "total_rows": len(rows) + len(errors),
                "imported_rows": len(rows),
                "created_rows": created,
                "updated_rows": updated,
                "error_rows": errors,
            }
        finally:
            if owns_session:
                await db.close()

    async def backfill_zero_prices(
        self,
        *,
        session: Optional[AsyncSession] = None,
    ) -> Dict[str, Any]:
        owns_session = session is None
        db = session or SessionLocal()
        try:
            result = await db.execute(select(FakeCerpProduct).where(FakeCerpProduct.amount <= 0))
            products = list(result.scalars().all())
            now = datetime.now(timezone.utc)
            for product in products:
                product.amount = self._resolve_list_price(
                    product.amount,
                    code=product.code,
                    name=product.name,
                    spec1=product.spec1 or "",
                )
                product.updated_at = now
            if owns_session:
                await db.commit()
            return {"backfilled_price_rows": len(products)}
        finally:
            if owns_session:
                await db.close()

    async def ensure_seeded(self) -> Dict[str, Any]:
        async with SessionLocal() as session:
            result = await session.execute(select(func.count(FakeCerpProduct.id)))
            count = int(result.scalar() or 0)
            if count > 0:
                backfill_result = await self.backfill_zero_prices(session=session)
                camping_count_result = await session.execute(
                    select(func.count(FakeCerpProduct.id)).where(FakeCerpProduct.code.like("CAMP%"))
                )
                camping_count = int(camping_count_result.scalar() or 0)
                camping_result: Dict[str, Any] = {}
                fixture_path = self.resolve_camping_fixture_path()
                fixture_payload: List[Dict[str, Any]] = []
                if fixture_path.exists():
                    fixture_payload = json.loads(fixture_path.read_text(encoding="utf-8"))
                if camping_count < len(fixture_payload):
                    camping_result = await self.import_json_rows(fixture_payload, session=session)
                    camping_count = int(camping_result.get("imported_rows") or camping_count)
                if backfill_result.get("backfilled_price_rows") or camping_result:
                    await session.commit()
                return {
                    "seeded": bool(camping_result),
                    "product_count": count + int(camping_result.get("created_rows") or 0),
                    "camping_product_count": camping_count,
                    "supplemental_camping_seed": bool(camping_result),
                    **backfill_result,
                }
            path = self.resolve_default_workbook_path()
            if not path.exists():
                return {
                    "seeded": False,
                    "product_count": 0,
                    "missing_default_path": str(path),
                }
            result_payload = await self.import_workbook(path.read_bytes(), session=session)
            fixture_path = self.resolve_camping_fixture_path()
            camping_result: Dict[str, Any] = {}
            if fixture_path.exists():
                camping_result = await self.import_json_rows(
                    json.loads(fixture_path.read_text(encoding="utf-8")),
                    session=session,
                )
            await session.commit()
            return {
                "seeded": True,
                **result_payload,
                "camping_product_count": int(camping_result.get("imported_rows") or 0),
                "default_path": str(path),
            }

    async def status(self) -> Dict[str, Any]:
        async with SessionLocal() as session:
            count_result = await session.execute(select(func.count(FakeCerpProduct.id)))
            updated_result = await session.execute(select(func.max(FakeCerpProduct.imported_at)))
            return {
                "product_count": int(count_result.scalar() or 0),
                "last_import_at": updated_result.scalar(),
                "default_seed_path": str(self.resolve_default_workbook_path()),
            }
