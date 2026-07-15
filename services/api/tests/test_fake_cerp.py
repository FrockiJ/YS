from __future__ import annotations

import json
from collections import Counter
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from src.cerpModule.fake_client import FakeCERPClient
from src.services.cerp_service import CerpService
from src.services.fake_cerp_import_service import FakeCerpImportService
from src.services.fake_cerp_pricing import resolve_fake_list_price, resolve_fake_vip_price


PROJECT_ROOT = Path(__file__).resolve().parents[3]
CAMPING_FIXTURE_PATH = PROJECT_ROOT / "cerp" / "fake_camping_products.json"
DEFAULT_SEED_PATH = next((PROJECT_ROOT / "cerp").glob("*.xlsx"))


def _workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_fake_cerp_deterministic_stock_is_stable():
    first = FakeCerpImportService.deterministic_stock("P355300100")
    second = FakeCerpImportService.deterministic_stock("p355300100")
    assert first == second
    assert 0 <= first <= 50


def test_fake_cerp_import_parser_accepts_default_headers_and_stock_override():
    payload = _workbook_bytes(
        [
            ["code", "name", "spec1", "amount", "stock"],
            ["P355300100", "G-Soul Upgrade PE8 150m #0.6", "#0.6", 1200, 7],
            ["", "Missing Code", "#0.8", 1000, 3],
        ]
    )

    rows, errors = FakeCerpImportService.parse_workbook(payload)

    assert len(rows) == 1
    assert rows[0]["code"] == "P355300100"
    assert rows[0]["name"] == "G-Soul Upgrade PE8 150m #0.6"
    assert rows[0]["amount"] == Decimal("1200")
    assert rows[0]["stock"] == 7
    assert len(errors) == 1
    assert errors[0]["reason_code"] == "missing_required_field"


def test_fake_cerp_import_parser_generates_price_when_amount_is_zero():
    payload = _workbook_bytes(
        [
            ["code", "name", "spec1", "amount", "stock"],
            ["P355300100", "G-Soul Upgrade PE8 150m #0.6", "#0.6", 0, 7],
        ]
    )

    rows, errors = FakeCerpImportService.parse_workbook(payload)

    assert not errors
    assert rows[0]["amount"] > 0
    assert rows[0]["amount"] == resolve_fake_list_price(
        0,
        code="P355300100",
        name="G-Soul Upgrade PE8 150m #0.6",
        spec1="#0.6",
    )


def test_fake_cerp_json_parser_accepts_camping_metadata():
    payload = [
        {
            "code": "camp0001",
            "name": "TrailForge Alpine Dome Tent 01",
            "spec1": "2P / aluminum pole",
            "amount": 0,
            "stock": 12,
            "brand": "TrailForge",
            "supplier": "TrailForge",
            "category": "Tent & Shelter",
            "type": "Tent",
            "material": "210T ripstop polyester",
            "size": "2P",
            "color": "Forest Green",
            "photo_url": "https://example.test/tent.jpg",
        }
    ]

    rows, errors = FakeCerpImportService.parse_json_rows(payload)

    assert not errors
    assert rows[0]["code"] == "CAMP0001"
    assert rows[0]["amount"] > 0
    assert rows[0]["stock"] == 12
    assert rows[0]["raw_row"]["brand"] == "TrailForge"
    assert rows[0]["raw_row"]["supplier"] == "TrailForge"
    assert rows[0]["raw_row"]["photo_url"] == "https://example.test/tent.jpg"


def test_fake_cerp_pricing_is_stable_and_vip_is_85_percent_rounded_to_10():
    first = resolve_fake_list_price(0, code="P355300100", name="G-Soul Upgrade PE8 150m #0.6", spec1="#0.6")
    second = resolve_fake_list_price(0, code="P355300100", name="G-Soul Upgrade PE8 150m #0.6", spec1="#0.6")

    assert first == second
    assert Decimal("380") <= first <= Decimal("1280")
    assert resolve_fake_vip_price(Decimal("1200")) == Decimal("1020")


def test_fake_cerp_pricing_preserves_nonzero_import_amount():
    assert resolve_fake_list_price(1234, code="P355300100", name="G-Soul Upgrade PE8 150m #0.6") == Decimal("1234")


def test_fake_cerp_client_maps_legacy_response_shape():
    product = SimpleNamespace(
        code="P355300100",
        name="G-Soul Upgrade PE8 150m #0.6",
        spec1="#0.6",
        amount=Decimal("1200"),
        stock=9,
        raw_row={
            "brand": "YGK",
            "supplier": "YGK",
            "category": "Fishing Gear",
            "photo_url": "https://example.test/product.jpg",
        },
    )

    row = FakeCERPClient._row(product, include_warehouses=True)

    assert row["invn002"] == "P355300100"
    assert row["invn005"] == "G-Soul Upgrade PE8 150m #0.6"
    assert row["invn006"] == "YGK"
    assert row["brand"] == "YGK"
    assert row["supplier"] == "YGK"
    assert row["photo_url"] == "https://example.test/product.jpg"
    assert row["invn013"] == 1200.0
    assert row["price"] == 1200.0
    assert row["list_price"] == 1200.0
    assert row["invn015"] == 1020.0
    assert row["vip_price"] == 1020.0
    assert row["quote_price"] == 1020.0
    assert row["stock"] == 9
    assert row["wd4inv1as"][0]["inv1015"] == 9


def test_fake_cerp_client_filters_paramchar_by_code_name_and_spec():
    product = SimpleNamespace(
        code="P355300100",
        name="G-Soul Upgrade PE8 150m #0.6",
        spec1="#0.6",
        amount=Decimal("1200"),
        stock=9,
        raw_row={"brand": "YGK", "supplier": "YGK", "category": "Fishing Gear"},
    )

    assert FakeCERPClient._filter_products([product], paramchar1={"invn002": ["P355"]}) == [product]
    assert FakeCERPClient._filter_products([product], paramchar1={"invn005": ["upgrade"]}) == [product]
    assert FakeCERPClient._filter_products([product], paramchar1={"invn051": ["0.6"]}) == [product]
    assert FakeCERPClient._filter_products([product], paramchar1={"invn006": ["ygk"]}) == [product]
    assert FakeCERPClient._filter_products([product], paramchar1={"invn006": ["missing"]}) == []


def test_fake_cerp_does_not_use_supplier_as_brand():
    product = SimpleNamespace(
        code="P1",
        name="Unbranded Fishing Rod",
        spec1="7 ft",
        amount=Decimal("1200"),
        stock=3,
        raw_row={"supplier": "YS 釣竿供應", "category": "釣竿"},
    )

    row = FakeCERPClient._row(product, include_warehouses=False)

    assert row["brand"] == ""
    assert row["invn006"] == ""
    assert row["supplier"] == "YS 釣竿供應"


def test_fake_cerp_reclassifies_false_rod_accessories():
    for name in ("Rod Stand", "Rod Holder", "Rod Wax", "Rod T-Shirt"):
        product = SimpleNamespace(
            code=f"P-{name}",
            name=name,
            spec1="",
            amount=Decimal("100"),
            stock=1,
            raw_row={"supplier": "YS 釣竿供應", "category": "釣竿"},
        )

        row = FakeCERPClient._row(product, include_warehouses=False)

        assert row["category"] == "配件"
        assert FakeCERPClient._filter_products([product], paramchar1={"invn030": ["釣竿"]}) == []


def test_fake_camping_fixture_has_900_unique_products_with_complete_metadata():
    products = json.loads(CAMPING_FIXTURE_PATH.read_text(encoding="utf-8"))
    codes = [product["code"] for product in products]

    assert len(products) == 900
    assert codes[0] == "CAMP0001"
    assert codes[-1] == "CAMP0900"
    assert len(set(codes)) == 900
    assert all(product["brand"] == product["supplier"] for product in products)
    assert len({product["brand"] for product in products}) == 15
    assert Counter(product["category"] for product in products) == {category: 60 for category in {product["category"] for product in products}}
    assert all(product.get("category") and product.get("supplier") and product.get("photo_url") for product in products)


def test_fake_camping_fixture_rows_parse_with_generated_prices():
    products = json.loads(CAMPING_FIXTURE_PATH.read_text(encoding="utf-8"))
    rows, errors = FakeCerpImportService.parse_json_rows(products)

    assert not errors
    assert len(rows) == 900
    assert all(row["amount"] > 0 for row in rows)
    assert all(row["raw_row"].get("brand") for row in rows)
    assert all(row["raw_row"].get("category") and row["raw_row"].get("supplier") and row["raw_row"].get("photo_url") for row in rows)


def test_default_fake_cerp_seed_has_complete_metadata():
    workbook = load_workbook(DEFAULT_SEED_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [str(value or "").strip().lower() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    indexes = {header: headers.index(header) for header in ("category", "supplier", "photo_url")}
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(rows) == 1933
    assert all(str(row[indexes["category"]] or "").strip() for row in rows)
    assert all(str(row[indexes["supplier"]] or "").strip() for row in rows)
    assert all(str(row[indexes["photo_url"]] or "").strip() for row in rows)


def test_product_mapping_preserves_seed_photo_url():
    row = {
        "invn002": "P355300100",
        "invn005": "Seed product",
        "invn013": 1200,
        "invn015": 1020,
        "photo_url": "https://example.test/seed-product.jpg",
    }

    product = CerpService.map_row_to_product(row, [])

    assert product["photo_url"] == "https://example.test/seed-product.jpg"
